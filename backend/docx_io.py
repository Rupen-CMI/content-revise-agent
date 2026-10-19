import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO

from urllib.parse import urlparse

import docx
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement

from checks import _BULLET_LINE_RE
from config import CMI_HOST

log = logging.getLogger(__name__)

OUTPUT_DIR = Path(__file__).parent / "output"


@dataclass
class Link:
    anchor: str
    url: str

    @property
    def is_external(self) -> bool:
        host = urlparse(self.url).netloc.lower()
        return bool(host) and not host.endswith(CMI_HOST)


@dataclass
class Paragraph:
    index: int
    text: str
    links: list[Link] = field(default_factory=list)


@dataclass
class SourceArticle:
    title: str
    paragraphs: list[Paragraph]

    @property
    def body(self) -> str:
        return "\n\n".join(p.text for p in self.paragraphs)

    def sponsor_candidates(self) -> list[Paragraph]:
        """Paragraphs carrying at least one non-CMI link; the paid backlink is one of these."""
        return [p for p in self.paragraphs if any(link.is_external for link in p.links)]


def read_docx(file: str | BinaryIO) -> SourceArticle:
    try:
        document = docx.Document(file)
    except Exception as e:
        log.error(f"could not open the uploaded file as a .docx: {type(e).__name__}: {e}", exc_info=True)
        raise ValueError(
            "This file could not be read as a Word document (.docx). "
            "If it came from Google Docs or an older .doc file, re-save it as .docx and try again."
        ) from e

    rels = {rid: rel.target_ref for rid, rel in document.part.rels.items() if "hyperlink" in rel.reltype}

    paragraphs: list[Paragraph] = []
    for p in document.paragraphs:
        text = "".join(t.text or "" for t in p._p.iter(qn("w:t"))).strip()
        if not text:
            continue
        links = []
        for h in p._p.findall(qn("w:hyperlink")):
            url = rels.get(h.get(qn("r:id")))
            if url:
                anchor = "".join(t.text or "" for t in h.iter(qn("w:t")))
                links.append(Link(anchor=anchor, url=url))
        paragraphs.append(Paragraph(index=len(paragraphs), text=text, links=links))

    if not paragraphs:
        raise ValueError("document has no text")
    title, *body = paragraphs
    for i, p in enumerate(body):
        p.index = i
    return SourceArticle(title=title.text, paragraphs=body)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    """python-docx has no native hyperlink support; this builds the required w:hyperlink
    XML by hand, the standard workaround for this library."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)

    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def output_path_for(original_filename: str) -> Path:
    OUTPUT_DIR.mkdir(exist_ok=True)
    stem = Path(original_filename).stem
    return OUTPUT_DIR / f"Revised_{stem}.docx"


# A paragraph that is a short lead-in sentence followed by parallel items (3 trends, a
# segmentation breakdown, ...) is written as ONE paragraph string with the items as "\n- " lines,
# rather than as separate paragraph objects — this keeps every existing check (word counts,
# keyword matching, verbatim overlap, citation lookup by exact string) working unchanged, since
# they all operate on plain text and treat "\n" like any other whitespace (see checks.py's
# count_items_in for the one check that needs to count bullets as separate items). Only the
# renderers (here and app.js's renderSections) need to know about the convention, splitting it
# back into a lead-in line plus real bullet-point paragraphs for display.


def _split_bullet_lines(para_text: str) -> tuple[str, list[str]]:
    """(lead-in text, [bullet item, ...]) if para_text contains "\n- " bullet lines, else
    (para_text, [])."""
    lines = para_text.split("\n")
    lead = []
    items = []
    for line in lines:
        if _BULLET_LINE_RE.match(line):
            items.append(_BULLET_LINE_RE.sub("", line, count=1))
        elif items:
            # a non-bullet line after bullets started: treat the whole thing as plain text
            # instead, rather than silently dropping a line that didn't fit the pattern.
            return para_text, []
        else:
            lead.append(line)
    return "\n".join(lead), items


def write_docx(
    path: Path, title: str, sections: list[dict], segments_fn,
) -> None:
    """Writes the revamped article as a .docx: each paragraph's text segments (from
    segments_fn, which expands the sponsor/CMI-link placeholders) become real hyperlink
    runs where a URL is present, plain runs otherwise. A paragraph containing "\n- " lines (see
    _split_bullet_lines) renders as a lead-in paragraph followed by a real bulleted list."""
    document = docx.Document()
    document.add_heading(title, level=1)

    def _add_run_segments(p, text: str) -> None:
        for seg_text, url in segments_fn(text):
            if url:
                _add_hyperlink(p, seg_text, url)
            else:
                p.add_run(seg_text)

    for s in sections:
        if s["heading"]:
            document.add_heading(s["heading"], level=2)
        for para_text in s["paragraphs"]:
            lead, items = _split_bullet_lines(para_text)
            if lead.strip():
                _add_run_segments(document.add_paragraph(), lead)
            for item in items:
                _add_run_segments(document.add_paragraph(style="List Bullet"), item)
            if not items and not lead.strip():
                document.add_paragraph()  # preserve an intentionally blank paragraph, if any

    OUTPUT_DIR.mkdir(exist_ok=True)
    document.save(str(path))
    log.info(f"wrote revamped article to {path}")
