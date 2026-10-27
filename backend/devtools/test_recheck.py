"""Offline test for mapping a humanized .docx back onto its run (no network, no LLM).

The round trip is the main check: a draft written to .docx the way the app writes it, read back with
recheck.parse, must pass and fail exactly the same SOP checks as the original draft.

Run from backend/: python devtools/test_recheck.py
"""
import contextlib
import copy
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import docx  # noqa: E402

import checks  # noqa: E402
import pipeline  # noqa: E402
import prompts  # noqa: E402
import recheck  # noqa: E402
from docx_io import Link, Paragraph, write_docx  # noqa: E402

with contextlib.redirect_stdout(io.StringIO()):
    import test_checks_v2 as t  # noqa: E402  (its own tests run on import; their output is not this test's)

FAILURES: list[str] = []
REPORT_URL = "https://www.coherentmarketinsights.com/market-insight/zzz-widget-services-market-1"
SPONSOR = Paragraph(
    index=0, text=t.SPONSOR_TEXT,
    links=[Link(anchor="Northway Specialists", url="https://northway-zz.com/widgets")])
OUT = Path(__file__).resolve().parent / "_recheck_tmp.docx"


def check(name: str, ok: bool, detail: str = "") -> None:
    if not ok:
        FAILURES.append(f"{name}: {detail}")


def written(sections: list[dict], faq: list[dict]) -> bytes:
    write_docx(OUT, t.TITLE, sections + pipeline.faq_sections(faq),
               lambda text: pipeline.expand_segments(text, SPONSOR, REPORT_URL))
    raw = OUT.read_bytes()
    OUT.unlink()
    return raw


def verdicts(sections, faq, title=t.TITLE) -> dict:
    return {r.name: r.passed for r in t.run_v2(title, sections, faq=faq)}


original = t.good_draft()
raw = written(original, t.GOOD_FAQ)
parsed = recheck.parse(raw, t.TITLE, original, t.GOOD_FAQ, t.SPONSOR_TEXT)

check("round trip: title", parsed["title"] == t.TITLE, parsed["title"])
check("round trip: no mapping notes", parsed["notes"] == [], str(parsed["notes"]))
check("round trip: same section tags in the same order",
      [s["sop"] for s in parsed["sections"]] == [s["sop"] for s in original])
check("round trip: same headings", [s["heading"] for s in parsed["sections"]] == [s["heading"] for s in original],
      str([s["heading"] for s in parsed["sections"]]))
check("round trip: the sponsor paragraph is back as its placeholder",
      sum(p.count(prompts.SPONSOR_TOKEN) for s in parsed["sections"] for p in s["paragraphs"]) == 1)
check("round trip: the CMI link is back as its placeholder",
      any("[[CMI_LINK:" in p for p in parsed["sections"][0]["paragraphs"]))
check("round trip: bullets rejoin their lead-in",
      any("\n- " in p for s in parsed["sections"] for p in s["paragraphs"]))
check("round trip: the FAQ keeps its entries and ids",
      [q.get("id") for q in parsed["faq"]] == [q.get("id") for q in t.GOOD_FAQ], str(parsed["faq"]))
check("round trip: opening/transition split lands on the original boundary",
      parsed["sections"][0]["paragraphs"] == original[0]["paragraphs"]
      and parsed["sections"][1]["paragraphs"] == original[1]["paragraphs"],
      str([len(s["paragraphs"]) for s in parsed["sections"][:2]]))
before = verdicts(original, t.GOOD_FAQ)
after = verdicts(parsed["sections"], parsed["faq"], parsed["title"])
check("round trip: the SOP checks give the same verdicts", before == after,
      str({k: (before.get(k), after.get(k)) for k in set(before) | set(after) if before.get(k) != after.get(k)}))

# a section's text edited by the humanizer: still mapped, and the check layer sees the edit
edited = copy.deepcopy(original)
edited[3]["paragraphs"][0] = "A fully rewritten lead-in sentence about the trends."
parsed2 = recheck.parse(written(edited, t.GOOD_FAQ), t.TITLE, original, t.GOOD_FAQ, t.SPONSOR_TEXT)
check("edited: the edit is carried through", parsed2["sections"][3]["paragraphs"][0].startswith("A fully rewritten"))
check("edited: nothing else moves", [s["heading"] for s in parsed2["sections"]] == [s["heading"] for s in original])

# a sponsor paragraph the humanizer removed is reported and never invented back
no_sponsor = copy.deepcopy(original)
no_sponsor[0]["paragraphs"][1] = ""
parsed3 = recheck.parse(written(no_sponsor, t.GOOD_FAQ), t.TITLE, original, t.GOOD_FAQ, t.SPONSOR_TEXT)
check("missing sponsor: a note says so", any("sponsor" in n for n in parsed3["notes"]), str(parsed3["notes"]))
check("missing sponsor: no placeholder is invented",
      all(prompts.SPONSOR_TOKEN not in q for s in parsed3["sections"] for q in s["paragraphs"]))

# a heading the humanizer dropped is reported, and its section comes back empty
no_head = copy.deepcopy(original)
no_head[5]["heading"] = ""
parsed4 = recheck.parse(written(no_head, t.GOOD_FAQ), t.TITLE, original, t.GOOD_FAQ, t.SPONSOR_TEXT)
check("dropped heading: reported", any(original[5]["heading"] in n for n in parsed4["notes"]), str(parsed4["notes"]))
check("dropped heading: the section's heading stays empty, so the structure checks can fail on it",
      parsed4["sections"][5]["heading"] == "", parsed4["sections"][5]["heading"])
check("dropped heading: the draft-structure SOP check fails",
      not next(v for k, v in verdicts(parsed4["sections"], parsed4["faq"], parsed4["title"]).items()
               if k.startswith("Draft structure matches")))

# headings that lost their Word style are still found by their wording
raw5 = written(original, t.GOOD_FAQ)
doc5 = docx.Document(io.BytesIO(raw5))
for para in doc5.paragraphs:
    if para.style.name == "Heading 2":
        para.style = doc5.styles["Normal"]
buf = io.BytesIO()
doc5.save(buf)
parsed5 = recheck.parse(buf.getvalue(), t.TITLE, original, t.GOOD_FAQ, t.SPONSOR_TEXT)
check("unstyled headings: still matched by wording", parsed5["notes"] == [], str(parsed5["notes"]))
check("unstyled headings: the sections are intact",
      [s["heading"] for s in parsed5["sections"]] == [s["heading"] for s in original])


# an editor changes the title, every heading and the FAQ questions: they are paired by position, and the
# humanized checks no longer judge the plan's own identity (title/heading/question wording, a keyword's section)
rewritten = copy.deepcopy(original)
n = 0
for sec in rewritten:
    if sec["heading"]:
        n += 1
        sec["heading"] = f"Entirely rewritten heading number {n} for readers" + ("?" if sec["heading"].endswith("?") else "")
faq2 = [{**q, "question": f"A freshly worded question number {i}?"} for i, q in enumerate(t.GOOD_FAQ)]
new_title = f"Why The {t.MARKET} Keeps Changing Faster Each Year"
raw6 = written(rewritten, faq2)
doc6 = docx.Document(io.BytesIO(raw6))
doc6.paragraphs[0].text = new_title
buf6 = io.BytesIO()
doc6.save(buf6)
parsed6 = recheck.parse(buf6.getvalue(), t.TITLE, original, t.GOOD_FAQ, t.SPONSOR_TEXT)
check("reworded headings: paired by position, no notes", parsed6["notes"] == [], str(parsed6["notes"]))
check("reworded headings: every section keeps its tag and carries the new heading",
      [(s["sop"], s["heading"]) for s in parsed6["sections"]] == [(s["sop"], s["heading"]) for s in rewritten],
      str([(s["sop"], s["heading"]) for s in parsed6["sections"]]))
check("reworded questions: the FAQ entries keep their ids", [q.get("id") for q in parsed6["faq"]] == [q.get("id") for q in t.GOOD_FAQ],
      str(parsed6["faq"]))
humanized = t.run_v2(parsed6["title"], parsed6["sections"], faq=parsed6["faq"], humanized=True)
names = {r.name: r.passed for r in humanized}
check("humanized: the plan-identity checks are not run",
      not any(k.startswith(("Draft structure matches", "Headings: question headings are planned")) for k in names), str(list(names)))
check("humanized: the SOP rule behind them is checked instead", names.get("2-3 headings are questions (most searched questions)") is True,
      str(names.get("2-3 headings are questions (most searched questions)")))
check("humanized: a reworded FAQ question is not a failure", all(
    "planned MSQ" not in r.detail for r in humanized), str([r.detail for r in humanized if "planned MSQ" in r.detail]))
check("humanized: keyword check names the article, not a section",
      any(k.startswith("Secondary keywords: 4-6 assigned, each used in the article") for k in names))
check("humanized: nothing fails because of the edits", all(names.values()), str([k for k, v in names.items() if not v]))
default_names = {r.name for r in t.run_v2(t.TITLE, original, faq=t.GOOD_FAQ)}
check("final screen keeps the plan-identity checks", any(k.startswith("Draft structure matches") for k in default_names))

if FAILURES:
    print("FAIL")
    for f in FAILURES:
        print(" -", f)
    sys.exit(1)
print("all recheck tests passed")
