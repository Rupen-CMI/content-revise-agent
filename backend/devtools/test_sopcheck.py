"""Offline test for the standalone SOP compliance check of a finished draft (no network, no LLM).

Run from backend/: python devtools/test_sopcheck.py
"""
import sys
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx  # noqa: E402

import sopcheck  # noqa: E402
from docx_io import _add_hyperlink  # noqa: E402

FAILURES: list[str] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    if not ok:
        FAILURES.append(f"{name}: {detail}")


def build(headed: bool = True, cmi_link: bool = True, outside_link: bool = True) -> bytes:
    d = docx.Document()
    d.add_heading("A Plain Title", level=1)
    p = d.add_paragraph("A short opening about widgets. ")
    if outside_link:
        _add_hyperlink(p, "Acme Widgets", "https://acme-zz.com/widgets")
    p2 = d.add_paragraph("The Zzz Widget Services Market is covered in ")
    if cmi_link:
        _add_hyperlink(p2, "Zzz Widget Services Market", "https://www.coherentmarketinsights.com/market-insight/zzz")
    head = (lambda t: d.add_heading(t, level=2)) if headed else (lambda t: d.add_paragraph(t))
    head("What Is The Size Of The Zzz Market?")
    d.add_paragraph("The market is growing. These are the key players in the market. It shows steady demand.")
    head("How Do Buyers Choose A Zzz Provider?")
    d.add_paragraph("Buyers compare price and support.")
    head("Frequently Asked Questions")
    d.add_paragraph("Who leads the market?")
    d.add_paragraph("Several providers lead the market.")
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def by_name(result: dict, fragment: str) -> dict:
    return next(c for c in result["checks"] if fragment in c["name"])


good = sopcheck.check_draft(build())
check("links: cmi once", by_name(good, "CMI report linked exactly once")["passed"])
check("links: cmi early", by_name(good, "CMI link sits within")["passed"])
check("links: sponsor in opening", by_name(good, "outside backlink sits in the opening")["passed"])
check("headings: 2 questions", by_name(good, "headings are questions")["passed"], str(by_name(good, "headings are questions")))
check("faq present", by_name(good, "FAQ section present")["passed"])
check("pointing-word check is disconnected", all("names what it refers to" not in c["name"] for c in good["checks"]))
check("length is soft", by_name(good, "Length is")["soft"] and not by_name(good, "Length is")["passed"])
check("skipped list is reported", len(good["skipped"]) >= 5)
check("title read", good["title"] == "A Plain Title", good["title"])
check("faq heading is not a section heading", "FAQ" not in by_name(good, "headings are questions")["detail"])

no_cmi = sopcheck.check_draft(build(cmi_link=False, outside_link=False))
check("no cmi link fails", not by_name(no_cmi, "CMI report linked exactly once")["passed"])
check("no sponsor link fails", not by_name(no_cmi, "outside backlink")["passed"])

flat = sopcheck.check_draft(build(headed=False))
check("no headings: heading checks skipped, not failed",
      all("headings are questions" not in c["name"] for c in flat["checks"])
      and any("headings are questions" in s["name"] for s in flat["skipped"]))

if FAILURES:
    print("FAIL")
    for f in FAILURES:
        print(" -", f)
    sys.exit(1)
print("ok: sopcheck")
