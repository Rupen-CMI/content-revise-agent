"""Maps a humanized .docx back onto the run that produced it, so the run's own SOP checks and keyword
highlighting can be applied to the edited draft.

A run's draft is a list of sections ({"heading", "sop", "paragraphs"}): the opening (untagged), the
transition, the market sections, plus the FAQ as separate question/answer pairs. A humanized draft is
only paragraphs and headings, so this module rebuilds that shape from the run's own headings: the
sponsor paragraph and the CMI link go back to their placeholders, bullet lines rejoin the paragraph
they belong to, and the opening splits from the transition by the original draft's proportions. Where
the mapping cannot be trusted a note says so, and the check that depends on it fails honestly instead
of passing by accident.
"""
import difflib
import re
from io import BytesIO
from urllib.parse import urlparse

import docx
from docx.oxml.ns import qn

import checks
from config import CMI_HOST
from prompts import SPONSOR_TOKEN

FAQ_HEADING = "Frequently Asked Questions"
_BULLET_MARKERS = ("•", "-", "–", "*")
_HEADING_MIN_RATIO = 0.6        # a Word-styled heading must at least resemble a planned heading
_PLAIN_HEADING_MIN_RATIO = 0.85  # an unstyled line is a heading only when it is nearly a planned one
_QUESTION_MAX_WORDS = 30


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _ratio(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, _norm(a), _norm(b)).ratio()


def _is_cmi(url: str) -> bool:
    return urlparse(url).netloc.lower().endswith(CMI_HOST)


def read_blocks(raw: bytes) -> list[dict]:
    """Every non-empty paragraph of the docx as {text, heading, bullet, bold}. A hyperlink to the CMI
    site becomes the CMI link placeholder, the form the checks and the renderer expect."""
    document = docx.Document(BytesIO(raw))
    rels = {rid: rel.target_ref for rid, rel in document.part.rels.items() if "hyperlink" in rel.reltype}
    blocks: list[dict] = []
    for p in document.paragraphs:
        pieces: list[str] = []
        for child in p._p:
            text = "".join(t.text or "" for t in child.iter(qn("w:t")))
            if child.tag == qn("w:hyperlink"):
                url = rels.get(child.get(qn("r:id")))
                if url and _is_cmi(url) and text.strip():
                    text = f"[[CMI_LINK:{text}]]"
            pieces.append(text)
        text = "".join(pieces).strip()
        if not text:
            continue
        try:
            style = p.style.name or ""
        except Exception:
            style = ""
        bullet = style.startswith("List") or text.startswith(_BULLET_MARKERS) and len(text) > 2 and text[1] == " "
        if bullet and text.startswith(_BULLET_MARKERS):
            text = text[2:].strip()
        runs = [r for r in p.runs if r.text.strip()]
        blocks.append({
            "text": text,
            "heading": style.startswith("Heading") or style == "Title",
            "bullet": bool(bullet),
            "bold": bool(runs) and all(r.bold for r in runs),
        })
    return blocks


def _tokenize_sponsor(text: str, sponsor_text: str) -> tuple[str, str]:
    """(text with the sponsor paragraph swapped for its placeholder, status). The swap happens only on an
    exact (whitespace-normalised) match: a sponsor paragraph the humanizer reworded must fail the
    'placed once, unaltered' check, not be silently restored. status is 'exact', 'changed' or ''."""
    if not sponsor_text.strip():
        return text, ""
    flat, want = " ".join(text.split()), " ".join(sponsor_text.split())
    i = flat.find(want)
    if i >= 0:
        return (flat[:i] + SPONSOR_TOKEN + flat[i + len(want):]).strip(), "exact"
    match = difflib.SequenceMatcher(None, flat.lower(), want.lower(), autojunk=False).find_longest_match(
        0, len(flat), 0, len(want))
    return text, "changed" if match.size >= 0.6 * len(want) else ""


def _paragraphs(blocks: list[dict]) -> list[str]:
    """Paragraph strings in the run's own convention: bullet lines rejoin the paragraph before them as
    '\\n- item' lines (checks and renderers treat that as one paragraph with a bulleted list)."""
    out: list[str] = []
    for b in blocks:
        if b["bullet"]:
            if out:
                out[-1] += "\n- " + b["text"]
            else:
                out.append("- " + b["text"])
        else:
            out.append(b["text"])
    return out


def _words(paragraph: str, sponsor_text: str) -> int:
    return checks._words(checks._sponsor_replacement(paragraph, sponsor_text))


def _split_opening(pre: list[str], run_sections: list[dict], sponsor_text: str) -> tuple[list[str], list[str]]:
    """The unheaded paragraphs are the opening followed by the transition. The docx does not mark the
    boundary, so it falls where the original draft's opening/transition word proportion falls, never
    before the paragraph holding the sponsor (which belongs to the opening)."""
    orig_i = sum(_words(p, sponsor_text) for s in run_sections if not s.get("sop") for p in s["paragraphs"])
    orig_t = sum(_words(p, sponsor_text) for s in run_sections if s.get("sop") == ["transition"] for p in s["paragraphs"])
    if len(pre) < 2 or orig_t == 0:
        return pre, []
    first = max(1, 1 + max((i for i, p in enumerate(pre) if SPONSOR_TOKEN in p), default=-1))
    if first > len(pre) - 1:
        return pre, []
    target = orig_i / (orig_i + orig_t)
    counts = [_words(p, sponsor_text) for p in pre]
    total = sum(counts) or 1
    best, best_gap, running = first, 9.0, sum(counts[:first - 1])
    for k in range(first, len(pre)):
        running += counts[k - 1]
        gap = abs(running / total - target)
        if gap < best_gap:
            best, best_gap = k, gap
    return pre[:best], pre[best:]


def _parse_faq(blocks: list[dict], run_faq: list[dict]) -> list[dict]:
    """Question/answer pairs: a bold (or short '?') line starts a question, the paragraphs after it are
    its answer. A question is paired with the run's own FAQ entry it most resembles, so the entry's id
    carries over; the humanized wording is kept, so a reworded question is judged as written."""
    pairs: list[dict] = []
    for b in blocks:
        words = len(b["text"].split())
        is_question = (not b["bullet"]) and (b["bold"] or b["text"].rstrip().endswith("?")) and words <= _QUESTION_MAX_WORDS
        if is_question:
            pairs.append({"question": b["text"], "answer": ""})
        elif pairs:
            piece = ("\n- " + b["text"]) if b["bullet"] else b["text"]
            pairs[-1]["answer"] = (pairs[-1]["answer"] + ("" if b["bullet"] else " ") + piece).strip()
    free = list(run_faq)
    matched: dict[int, dict] = {}
    for i, qa in enumerate(pairs):
        if not qa["answer"] or not free:
            continue
        match = max(free, key=lambda f: _ratio(qa["question"], str(f.get("question") or "")))
        if _ratio(qa["question"], str(match.get("question") or "")) >= 0.6:
            free.remove(match)
            matched[i] = match
    # An editor may reword a question beyond recognition: when exactly as many questions are left as FAQ
    # entries, they are paired by position.
    left = [i for i, qa in enumerate(pairs) if qa["answer"] and i not in matched]
    if left and len(left) == len(free):
        for i, entry in zip(left, free):
            matched[i] = entry
    return [{**matched.get(i, {}), "question": qa["question"], "answer": qa["answer"]}
            for i, qa in enumerate(pairs) if qa["answer"]]


def parse(raw: bytes, run_title: str, run_sections: list[dict], run_faq: list[dict], sponsor_text: str) -> dict:
    """{title, sections, faq, notes}: the humanized draft in the run's own shape."""
    blocks = read_blocks(raw)
    if not blocks:
        raise ValueError("document has no text")
    notes: list[str] = []
    title = blocks[0]["text"]
    blocks = blocks[1:]

    wanted = [(i, s["heading"]) for i, s in enumerate(run_sections) if s.get("heading")]
    if run_faq:
        wanted.append((-1, FAQ_HEADING))
    unused = dict(wanted)  # run index -> planned heading

    # Which blocks are headings, and which planned heading each is. Every pairing is scored first and the best
    # pairs win globally, so two similar question headings cannot be crossed by the order they appear in; a
    # Word-styled heading no planned heading resembles is then paired by position, when exactly as many
    # styled headings and planned headings are left, because an editor may reword a heading freely.
    owner: dict[int, int] = {}   # block index -> run index (-1 = FAQ)
    pairs = []
    for bi, b in enumerate(blocks):
        if b["bullet"] or len(b["text"].split()) > 25:
            continue
        floor = _HEADING_MIN_RATIO if b["heading"] else _PLAIN_HEADING_MIN_RATIO
        for k, h in unused.items():
            r = _ratio(b["text"], h)
            if r >= floor:
                pairs.append((r, bi, k))
    for r, bi, k in sorted(pairs, key=lambda x: -x[0]):
        if bi not in owner and k in unused:
            owner[bi] = k
            del unused[k]
    left_styled = [bi for bi, b in enumerate(blocks) if b["heading"] and not b["bullet"] and bi not in owner]
    if left_styled and len(left_styled) == len(unused):
        for bi, k in zip(left_styled, sorted(unused, key=lambda k: (k == -1, k))):
            owner[bi] = k
        unused.clear()
        left_styled = []
    for bi in left_styled:
        notes.append(f"the heading {blocks[bi]['text']!r} matches none of this run's headings")
    for k, h in unused.items():
        notes.append(f"the heading {h!r} was not found in the humanized draft")

    buckets: dict[int, list[dict]] = {}
    current: int | None = None
    pre: list[dict] = []
    order: list[int] = []
    for bi, b in enumerate(blocks):
        if bi in owner:
            current = owner[bi]
            buckets[current] = []
            order.append(current)
            continue
        (pre if current is None else buckets[current]).append(b)
    if order != sorted(order, key=lambda k: (k == -1, k)):
        notes.append("the sections are in a different order than the run's draft")

    sponsor_status = ""

    def tokenize(blocks_: list[dict]) -> list[dict]:
        nonlocal sponsor_status
        out = []
        for b in blocks_:
            text, status = _tokenize_sponsor(b["text"], sponsor_text)
            sponsor_status = sponsor_status if sponsor_status == "exact" else (status or sponsor_status)
            out.append({**b, "text": text})
        return out

    pre_paragraphs = _paragraphs(tokenize(pre))
    opening, transition = _split_opening(pre_paragraphs, run_sections, sponsor_text)
    if sponsor_status == "changed":
        notes.append("the sponsor paragraph's wording was changed, so it cannot be treated as the retained paragraph")
    elif sponsor_text.strip() and not sponsor_status:
        notes.append("the sponsor paragraph was not found in the humanized draft")
    if transition == [] and any(s.get("sop") == ["transition"] for s in run_sections):
        notes.append("the opening and the transition could not be told apart (too few paragraphs)")

    sections: list[dict] = []
    heading_text = {bi: blocks[bi]["text"] for bi in owner}
    heading_of = {run_index: heading_text[bi] for bi, run_index in owner.items()}
    for i, s in enumerate(run_sections):
        if not s.get("sop"):
            paragraphs, heading = opening, ""
        elif s["sop"] == ["transition"]:
            paragraphs, heading = transition, ""
        else:
            paragraphs = _paragraphs(tokenize(buckets.get(i, [])))
            heading = heading_of.get(i, "")  # a heading that was not found stays empty, so the structure checks fail on it
        sections.append({"heading": heading, "sop": s.get("sop", []), "paragraphs": paragraphs})

    faq = _parse_faq(buckets[-1], run_faq) if -1 in buckets else []
    return {"title": title, "sections": sections, "faq": faq, "notes": notes}
